"""
Regenerates feature2_v2/data/*.json from the Elnesr & Alazba (2016)
supplementary workbook (data/sources/1-s2.0-S0168169916300989-mmc1.xlsx).

Nothing in data/ is hand-typed: every number there comes out of the
workbook through this script, so the provenance is auditable.

Outputs
-------
stations_ksa.json      FAOCLIM-2 sinusoidal fits (Tx, Tn, Ta, ET0) for the
                       stations that are physically inside Saudi Arabia.
                       The workbook labels 33 stations "Saudi Arabia", but
                       several sit in Kuwait, Bahrain or Jordan; they are
                       removed with a point-in-polygon test against the
                       repo's national border (data/geo/national_border).
crops_eln16.json       All crop/condition rows of the "Crops" sheet (122; the paper says 123).
spreadsheet_fixture.json
                       The inputs and the CACHED computed values of the
                       workbook's "Model" sheet (365 sowing days). The test
                       suite checks that elnesr_model.py reproduces them.

Run:  pip install openpyxl  &&  python feature2_v2/build_data.py
"""

import json
import warnings
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
XLSX = ROOT / "data" / "sources" / "1-s2.0-S0168169916300989-mmc1.xlsx"
BORDER = ROOT / "data" / "geo" / "national_border" / "SAU-geo.json"
OUT = HERE / "data"

warnings.filterwarnings("ignore", module="openpyxl")


def _point_in_polygon(lon, lat, ring):
    """Ray casting; ring is a list of [lon, lat]."""
    inside = False
    j = len(ring) - 1
    for i in range(len(ring)):
        xi, yi = ring[i][0], ring[i][1]
        xj, yj = ring[j][0], ring[j][1]
        if (yi > lat) != (yj > lat) and lon < (xj - xi) * (lat - yi) / (yj - yi) + xi:
            inside = not inside
        j = i
    return inside


def _dist_to_ring_deg(lon, lat, ring):
    """Minimum planar distance (degrees) from a point to a polygon ring."""
    best = 1e9
    for (x1, y1), (x2, y2) in zip(ring[:-1], ring[1:]):
        dx, dy = x2 - x1, y2 - y1
        t = 0.0 if dx == dy == 0 else max(0.0, min(1.0, ((lon - x1) * dx + (lat - y1) * dy) / (dx * dx + dy * dy)))
        best = min(best, ((lon - x1 - t * dx) ** 2 + (lat - y1 - t * dy) ** 2) ** 0.5)
    return best


# The border file is generalised; coastal towns (e.g. Al Wajh) fall a few km
# outside it. Stations within this distance of the border line are kept.
# Checked: every foreign station in the workbook's "Saudi Arabia" rows
# (Kuwait, Bahrain, Jordan) is further away than this.
COAST_TOLERANCE_DEG = 0.1


def _in_ksa(lat, lon, polygons):
    return any(_point_in_polygon(lon, lat, poly[0]) or _dist_to_ring_deg(lon, lat, poly[0]) < COAST_TOLERANCE_DEG
               for poly in polygons)


def _border_polygons():
    g = json.loads(BORDER.read_text())
    polys = []
    for f in g["features"]:
        geom = f["geometry"]
        if geom["type"] == "Polygon":
            polys.append(geom["coordinates"])
        else:
            polys.extend(geom["coordinates"])
    return polys


def _fit(row, i):
    return {"a": row[i], "rho": row[i + 1], "omega": row[i + 2], "phi": row[i + 3]}


def build_stations(wb):
    ws = wb["Stations"]
    rows = list(ws.iter_rows(values_only=True))
    polys = _border_polygons()
    kept, dropped = [], []
    for r in rows[1:]:
        if not r[18] or "Saudi" not in str(r[18]):
            continue
        rec = {
            "id": r[0], "wmo_or_faoclim_id": r[17], "country_label": r[18],
            "region": r[19], "name": r[20], "lat": r[22], "lon": r[23], "elevation_m": r[24],
            "Tx": _fit(r, 1), "Tn": _fit(r, 5), "Ta": _fit(r, 9), "ET0": _fit(r, 13),
        }
        (kept if _in_ksa(rec["lat"], rec["lon"], polys) else dropped).append(rec)
    return {
        "source": "Elnesr & Alazba (2016) Appendix A workbook, sheet 'Stations' "
                  "(sinusoidal fits of FAOCLIM-2 long-term monthly means; FAO 2001). "
                  "Model: X(j) = a + rho * sin(omega * j + phi), j = day of year.",
        "filter": "Kept only stations inside data/geo/national_border/SAU-geo.json.",
        "dropped_outside_border": [f"{d['name']} ({d['lat']}, {d['lon']})" for d in dropped],
        "stations": kept,
    }


CROP_COLS = ["number", "family", "crop", "plant_date", "region", "row", "kc_ini", "kc_mid", "kc_end",
             "height_m", "dur_ini", "dur_dev", "dur_mid", "dur_late", "dur_total", "t_max", "t_min",
             "t_base", "t_opt", "dur_therm", "dur_before_harvest", "heat_tol_pct", "hu_min_tab", "hu_max_tab"]


def build_crops(wb):
    ws = wb["Crops"]
    out = []
    for r in list(ws.iter_rows(values_only=True))[1:]:
        # skip blank rows and the sheet's trailing column-index row ("1, 3, 4, 5, ...")
        if not isinstance(r[0], (int, float)) or not isinstance(r[6], (int, float)) or not isinstance(r[2], str):
            continue
        rec = dict(zip(CROP_COLS, r[:24]))
        for k in ("family", "crop", "plant_date", "region"):
            rec[k] = str(rec[k]).strip()
        for k in CROP_COLS[6:]:
            rec[k] = round(float(rec[k]), 4)
        rec["remarks"] = [str(x).strip() for x in r[24:27] if x not in (None, "")]
        out.append(rec)
    out.sort(key=lambda c: c["number"])
    return {
        "source": "Elnesr & Alazba (2016) Appendix A workbook, sheet 'Crops'. Kc and stage lengths "
                  "from FAO-56 (Allen et al. 1998) Tables 11-12; cardinal temperatures and heat "
                  "tolerance from Alsadon (2002), Maynard & Hochmuth (2006) and others cited in the paper.",
        "columns": {
            "t_max / t_min": "crTmax / crTmin, maximum / minimum tolerable temperature (degC)",
            "t_base / t_opt": "crTbase / crTopt, base and optimum temperature (degC)",
            "dur_therm": "thermal duration used by the Model sheet as the season length (days)",
            "heat_tol_pct": "Htol, heat-unit tolerance above optimal (%), Eq. 8",
            "hu_min_tab / hu_max_tab": "(t_opt - t_base) * dur_total and (1 + Htol/100) * that",
        },
        "crops": out,
    }


def build_fixture(wb_values, wb_formulas):
    ws = wb_values["Model"]
    wf = wb_formulas["Model"]
    v = lambda ref: ws[ref].value

    def fit(row):
        return {"a": v(f"AM{row}"), "rho": v(f"AN{row}"), "omega": v(f"AO{row}"), "phi": v(f"AP{row}")}

    # The "user-modified" column (D / G) is what the model computes with.
    params = {
        "kc_ini": v("D13"), "kc_mid": v("D14"), "kc_end": v("D15"),
        "dur_ini": v("D17"), "dur_dev": v("D18"), "dur_mid": v("D19"), "dur_late": v("D20"),
        "dur_total": v("G15"), "dur": v("G16"), "t_max": v("D21"), "t_min": v("D22"),
        "t_base": v("G21"), "t_opt": v("G22"), "heat_tol_pct": v("D23"), "hu_weight": v("D24"),
        "kc_eq_sheet": v("D16"), "hu_min_sheet": v("D25"), "hu_max_sheet": v("G25"),
    }
    cols = {"I": "doy", "J": "tx", "K": "tn", "L": "ta", "M": "et0", "N": "gdd",
            "O": "et_sum", "P": "hu_sum", "Q": "doy_harvest", "R": "et_int", "S": "hu_int",
            "T": "flag_hu", "U": "flag_tmax", "V": "flag_tmin", "W": "flag_all",
            "X": "idx_et_sum", "Y": "idx_hu_sum", "Z": "idx_comb_sum",
            "AA": "idx_et_int", "AB": "idx_hu_int", "AC": "idx_comb_int"}
    rows = []
    for r in range(27, 27 + 365):
        rows.append({name: ws[f"{c}{r}"].value for c, name in cols.items()})
    # Excel stores these as day serials in its 1900 date system; openpyxl turns
    # them into datetimes. Serial = day of year here (Excel's fictitious
    # 29 Feb 1900 makes serial 90 display as "Mar-30" though it is DOY 90).
    fmt = lambda d: (d.date() - __import__("datetime").date(1899, 12, 30)).days if hasattr(d, "date") else d
    return {
        "source": "Cached values of the Model sheet as shipped in mmc1.xlsx (Excel computed them). "
                  "Note: Model!L27 contains '(1+#REF!)*...', so Excel falls back to Ta = (Tx+Tn)/2.",
        "station": {"name": v("D6"), "state": v("D4"), "country": v("D2"),
                    "Tx": fit(3), "Tn": fit(4), "Ta": fit(5), "ET0": fit(6)},
        "params": params,
        "formula_L27": wf["L27"].value,
        "rows": rows,
        "summary": {
            "best_day_text": v("AR18"), "window_units": "Excel day serial (= day of year)",
            "hu_windows": [[fmt(v("AT21")), fmt(v("AU21"))], [fmt(v("AV21")), fmt(v("AW21"))]],
            "hu_temp_windows": [[fmt(v("AT22")), fmt(v("AU22"))], [fmt(v("AV22")), fmt(v("AW22"))]],
            "n_hu_days": v("AD20"), "n_hu_temp_days": v("AJ20"),
            "max_comb_idx_hu_sum": v("AF18"), "max_comb_idx_hu_int": v("AI18"),
        },
    }


def main():
    OUT.mkdir(exist_ok=True)
    wb_v = openpyxl.load_workbook(XLSX, data_only=True)
    wb_f = openpyxl.load_workbook(XLSX, data_only=False)
    stations = build_stations(wb_v)
    crops = build_crops(wb_v)
    fixture = build_fixture(wb_v, wb_f)
    (OUT / "stations_ksa.json").write_text(json.dumps(stations, indent=1, ensure_ascii=False))
    (OUT / "crops_eln16.json").write_text(json.dumps(crops, indent=1, ensure_ascii=False))
    (OUT / "spreadsheet_fixture.json").write_text(json.dumps(fixture, indent=1, ensure_ascii=False, default=str))
    print(f"stations kept: {len(stations['stations'])}, dropped: {stations['dropped_outside_border']}")
    print(f"crop rows: {len(crops['crops'])}")
    print(f"fixture rows: {len(fixture['rows'])}, summary: {fixture['summary']}")


if __name__ == "__main__":
    main()
